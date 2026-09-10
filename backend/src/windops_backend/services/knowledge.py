from __future__ import annotations

import asyncio
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from docx import Document
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from windops_backend.agents.embeddings import (
    DeterministicTestEmbeddingProvider,
    EmbeddingProvider,
    LiteLLMEmbeddingProvider,
    aggregate_embedding_vectors,
    chunk_embedding_text,
    embed_texts_with_controls,
    pack_embedding_vector,
)
from windops_backend.config import Settings
from windops_backend.enums import Environment
from windops_backend.errors import ConflictError, NotFoundError
from windops_backend.models import KnowledgeDocument
from windops_backend.outbox import (
    OutboxLeaseLostError,
    assert_current_claim,
    claim_event,
    enqueue_knowledge_graph_projection,
    mark_event_failed,
    mark_event_succeeded,
)
from windops_backend.schemas import KnowledgeDocumentCreateRequest
from windops_backend.services.events import append_domain_event
from windops_backend.services.knowledge_access import KnowledgeDocumentScope
from windops_backend.storage import OutboxEvent

KNOWLEDGE_DOCUMENT_INDEX_REQUESTED = "knowledge.document.index.requested"
MAX_KNOWLEDGE_ARTIFACT_BYTES = 50 * 1024 * 1024
MAX_EXTRACTED_TEXT_CHARACTERS = 4_000_000
ALLOWED_KNOWLEDGE_CONTENT_TYPES = frozenset(
    {
        "text/plain",
        "text/markdown",
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
)


def _extract_knowledge_text(content: bytes, content_type: str) -> str:
    try:
        if content_type in {"text/plain", "text/markdown"}:
            text = content.decode("utf-8")
        elif content_type == "application/pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("encrypted PDF knowledge documents are not accepted")
            text = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
        elif content_type == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            document = Document(BytesIO(content))
            text = "\n".join(paragraph.text.strip() for paragraph in document.paragraphs)
        else:  # pragma: no cover - schema and object-store checks prevent this
            raise ValueError("knowledge document content type is not supported")
    except ValueError:
        raise
    except Exception as exc:
        # Parser/library errors are input failures, not server failures. Keep
        # implementation details out of the API response while retaining the
        # original exception for secured worker logs and diagnostics.
        raise ValueError("knowledge document could not be parsed") from exc
    normalized = "\n".join(line.rstrip() for line in text.replace("\x00", "").splitlines()).strip()
    if not normalized:
        raise ValueError("knowledge document extraction produced no text")
    if len(normalized) > MAX_EXTRACTED_TEXT_CHARACTERS:
        raise ValueError("knowledge document extracted text exceeds the indexing limit")
    return normalized


async def extract_knowledge_text(content: bytes, content_type: str) -> str:
    return await asyncio.to_thread(_extract_knowledge_text, content, content_type)


async def create_knowledge_document(
    session: AsyncSession,
    request: KnowledgeDocumentCreateRequest,
    *,
    extracted_text: str,
    content_size_bytes: int,
    subject: str,
    scope: KnowledgeDocumentScope | None = None,
) -> tuple[KnowledgeDocument, bool, str | None]:
    resolved_scope = scope or KnowledgeDocumentScope(
        tenant_id=request.tenant_id,
        wind_farm_id=request.wind_farm_id,
        turbine_id=request.turbine_id,
        data_scope=request.data_scope,
    )
    existing = await session.get(KnowledgeDocument, request.document_id)
    if existing is not None:
        if existing.artifact_sha256 == request.artifact_sha256.lower():
            if (
                existing.tenant_id != resolved_scope.tenant_id
                or existing.wind_farm_id != resolved_scope.wind_farm_id
                or existing.turbine_id != resolved_scope.turbine_id
                or existing.data_scope != resolved_scope.data_scope
            ):
                raise ConflictError(
                    f"knowledge document {request.document_id} already exists "
                    "with a different scope"
                )
            return existing, True, None
        raise ConflictError(
            f"knowledge document {request.document_id} already exists with different content"
        )
    document = KnowledgeDocument(
        id=request.document_id,
        tenant_id=resolved_scope.tenant_id,
        wind_farm_id=resolved_scope.wind_farm_id,
        turbine_id=resolved_scope.turbine_id,
        data_scope=resolved_scope.data_scope,
        title=request.title,
        document_type=request.document_type,
        body=extracted_text,
        citation_uri=f"windops://knowledge/documents/{request.document_id}",
        artifact_uri=request.artifact_uri,
        artifact_sha256=request.artifact_sha256.lower(),
        content_type=request.content_type,
        content_size_bytes=content_size_bytes,
        document_version=request.document_version,
        metadata_=request.metadata,
        ingestion_status="pending",
        vectorized=False,
        created_by=subject,
    )
    session.add(document)
    event = OutboxEvent(
        event_type=KNOWLEDGE_DOCUMENT_INDEX_REQUESTED,
        aggregate_type="knowledge_document",
        aggregate_id=document.id,
        payload={"document_id": document.id},
    )
    session.add(event)
    append_domain_event(
        session,
        event_type="knowledge.document.ingested",
        aggregate_type="knowledge_document",
        aggregate_id=document.id,
        payload={
            "document_id": document.id,
            "document_type": document.document_type,
            "document_version": document.document_version,
            "artifact_sha256": document.artifact_sha256,
            "content_size_bytes": document.content_size_bytes,
            "created_by": subject,
        },
    )
    await session.flush()
    return document, False, event.id


def _embedding_provider(settings: Settings) -> EmbeddingProvider:
    if settings.environment is Environment.TEST:
        return DeterministicTestEmbeddingProvider()
    return LiteLLMEmbeddingProvider(settings.embedding_model)


async def _mark_document_failed(
    factory: async_sessionmaker[AsyncSession], document_id: str
) -> None:
    async with factory() as session, session.begin():
        document = await session.get(KnowledgeDocument, document_id)
        if document is not None and not document.vectorized:
            document.ingestion_status = "failed"


async def process_knowledge_document_index_event(
    factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    event_id: str,
    provider: EmbeddingProvider | None = None,
) -> bool:
    async with factory() as claim_session:
        async with claim_session.begin():
            event = await claim_event(claim_session, event_id)
        if event is None:
            return False
        claim_token = event.claim_token
        if claim_token is None:  # pragma: no cover - protected by claim update
            raise RuntimeError("claimed knowledge index event has no fencing token")
        if event.event_type != KNOWLEDGE_DOCUMENT_INDEX_REQUESTED:
            error = ValueError(f"unexpected knowledge event type: {event.event_type}")
            await mark_event_failed(factory, event_id, claim_token, error)
            raise error
        document_id = str(event.payload["document_id"])
        try:
            async with factory() as read_session:
                document = await read_session.get(KnowledgeDocument, document_id)
                if document is None:
                    raise NotFoundError(f"knowledge document {document_id} was not found")
                if document.vectorized:
                    return await mark_event_succeeded(factory, event_id, claim_token)
                embedding_inputs = chunk_embedding_text(document.title, document.body)
            selected_provider = provider or _embedding_provider(settings)
            vectors = await embed_texts_with_controls(
                selected_provider,
                embedding_inputs,
                settings=settings,
            )
            vector = aggregate_embedding_vectors(vectors)
            async with factory() as write_session, write_session.begin():
                await assert_current_claim(write_session, event_id, claim_token)
                document = await write_session.scalar(
                    select(KnowledgeDocument)
                    .where(KnowledgeDocument.id == document_id)
                    .with_for_update()
                )
                if document is None:
                    raise NotFoundError(f"knowledge document {document_id} was not found")
                if not document.vectorized:
                    dialect = (
                        write_session.bind.dialect.name
                        if write_session.bind is not None
                        else "unknown"
                    )
                    document.embedding = (
                        pack_embedding_vector(vector) if dialect == "sqlite" else vector
                    )
                    document.vectorized = True
                    document.ingestion_status = "indexed"
                    document.embedding_provider = selected_provider.provider_name
                    document.embedding_model = selected_provider.model_name
                    document.indexed_at = datetime.now(UTC)
                    document.updated_at = document.indexed_at
                    append_domain_event(
                        write_session,
                        event_type="knowledge.document.indexed",
                        aggregate_type="knowledge_document",
                        aggregate_id=document.id,
                        payload={
                            "document_id": document.id,
                            "embedding_provider": document.embedding_provider,
                            "embedding_model": document.embedding_model,
                        },
                    )
                    enqueue_knowledge_graph_projection(
                        write_session,
                        aggregate_type="knowledge_document",
                        aggregate_id=document.id,
                        reason="knowledge-document-indexed",
                    )
        except asyncio.CancelledError:
            # Do not mark a document failed when the worker is cancelled. The
            # processing lease remains reclaimable and the document stays
            # pending for a later indexing attempt.
            raise
        except OutboxLeaseLostError:
            return False
        except BaseException as exc:
            await _mark_document_failed(factory, document_id)
            await mark_event_failed(factory, event_id, claim_token, exc)
            raise
    return await mark_event_succeeded(factory, event_id, claim_token)


def serialize_knowledge_document(document: KnowledgeDocument) -> dict[str, Any]:
    return {
        "document_id": document.id,
        "tenant_id": document.tenant_id,
        "wind_farm_id": document.wind_farm_id,
        "turbine_id": document.turbine_id,
        "data_scope": document.data_scope,
        "title": document.title,
        "document_type": document.document_type,
        "document_version": document.document_version,
        "citation_uri": document.citation_uri,
        "artifact_uri": document.artifact_uri,
        "artifact_sha256": document.artifact_sha256,
        "content_type": document.content_type,
        "content_size_bytes": document.content_size_bytes,
        "metadata": document.metadata_,
        "ingestion_status": document.ingestion_status,
        "vectorized": document.vectorized,
        "embedding_provider": document.embedding_provider,
        "embedding_model": document.embedding_model,
        "indexed_at": document.indexed_at,
        "created_by": document.created_by,
        "created_at": document.created_at,
        "updated_at": document.updated_at,
    }
