"""Native source-location and permission contracts; synthetic documents only."""

import hashlib
from io import BytesIO
from uuid import uuid4

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import select

from windops_backend.agents.embeddings import DeterministicTestEmbeddingProvider
from windops_backend.agents.tools import SQLToolAdapter
from windops_backend.knowledge_graph.domain import GraphAccessPolicy, NodeType
from windops_backend.knowledge_graph.projection import build_knowledge_graph_snapshot
from windops_backend.knowledge_parsing import (
    DOCX_TYPE,
    MAX_PASSAGES,
    PASSAGE_CHARACTERS,
    parse_knowledge_content,
    parse_legacy_body,
)
from windops_backend.models import KnowledgeDocument, KnowledgePassage, Tenant
from windops_backend.read_audit import ReadAuditEnqueueError
from windops_backend.schemas import KnowledgeDocumentCreateRequest
from windops_backend.services.knowledge import (
    create_knowledge_document,
    process_knowledge_document_index_event,
)
from windops_backend.storage import OutboxEvent


def pdf_with_second_page_text() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    page = writer.add_blank_page(width=300, height=400)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 200 Td (Synthetic tendon retest procedure) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


async def ingest(app, client, content, content_type="text/markdown", **changes):
    document_id = changes.pop("document_id", f"KB-PASSAGE-{uuid4().hex[:12]}")
    uri = f"minio://windops-knowledge-documents/documents/{document_id}/source"
    sha = app.state.artifact_verifier.register_object(uri, content, content_type)
    payload = {
        "document_id": document_id,
        "title": "Synthetic hybrid tower retest source",
        "document_type": "maintenance-procedure",
        "artifact_uri": uri,
        "artifact_sha256": sha,
        "content_type": content_type,
        "document_version": "synthetic-r1",
        **changes,
    }
    result = await client.post(
        "/api/v1/knowledge/documents",
        json=payload,
        headers={"Idempotency-Key": str(uuid4())},
    )
    assert result.status_code == 202, result.text
    return payload


async def test_pdf_positions_preserve_blank_page_and_original_source_hash(app, client):
    content = pdf_with_second_page_text()
    payload = await ingest(
        app,
        client,
        content,
        "application/pdf",
        metadata={
            "page_count": 999,
            "_openvigil_source_layout": {"page_count": 999},
        },
    )
    document_id = payload["document_id"]
    layout = (await client.get(f"/api/v1/knowledge/documents/{document_id}")).json()[
        "source_layout"
    ]
    assert layout["page_count"] == 2
    assert layout["unextracted_pages"] == [1]
    rows = (await client.get(f"/api/v1/knowledge/documents/{document_id}/passages")).json()[
        "passages"
    ]
    assert len(rows) == 1
    passage = rows[0]
    assert passage["page_number"] == 2
    assert passage["native_locator"]["page_number"] == 2
    assert passage["native_locator"]["ocr"] is False
    assert passage["native_locator"]["bbox"] is None
    assert passage["source_sha256"] == hashlib.sha256(content).hexdigest()
    assert passage["text_sha256"] == hashlib.sha256(passage["text"].encode()).hexdigest()
    assert passage["document_version"] == "synthetic-r1"
    detail = await client.get(passage["citation_href"])
    assert detail.status_code == 200
    assert detail.json()["passage_id"] == passage["passage_id"]
    source = await client.get(passage["source_href"])
    assert source.status_code == 200, source.text
    assert source.json()["artifact_sha256"] == payload["artifact_sha256"]
    assert source.json()["page_number"] == 2
    assert source.json()["download_url"].startswith("memory://")
    assert source.json()["expires_at"]
    # The source object cannot be substituted beneath its immutable document.
    app.state.artifact_verifier.register_object(
        payload["artifact_uri"], b"changed", "application/pdf"
    )
    rejected = await client.get(passage["source_href"])
    assert rejected.status_code == 422
    assert rejected.json()["error"]["code"] == "INVALID_TRANSITION"
    assert "SHA-256" in rejected.json()["error"]["message"]


async def test_markdown_sections_lines_and_ordinal_pagination(app, client):
    content = (
        b"# Calibration\n\nCheck synthetic sensor certificate.\n\n## Retest\n\n"
        b"Record direct force and temperature.\n"
    )
    payload = await ingest(app, client, content)
    document_id = payload["document_id"]
    base = f"/api/v1/knowledge/documents/{document_id}/passages"
    page = (await client.get(base, params={"limit": 1})).json()
    assert page["count"] == 1 and page["next_cursor"] == 0
    assert page["passages"][0]["native_locator"]["line_start"] == 1
    following = (await client.get(base, params={"cursor": 0, "limit": 100})).json()
    assert following["passages"][0]["ordinal"] == 1
    assert following["passages"][0]["native_locator"]["line_start"] == 3
    assert following["passages"][-1]["section"] == "Retest"
    body = (await client.get(f"/api/v1/knowledge/documents/{document_id}")).json()["body"]
    for row in [*page["passages"], *following["passages"]]:
        assert row["page_number"] is None
        assert body[row["char_start"] : row["char_end"]] == row["text"]
    changed_version = await client.post(
        "/api/v1/knowledge/documents",
        json={**payload, "document_version": "synthetic-r2"},
        headers={"Idempotency-Key": str(uuid4())},
    )
    assert changed_version.status_code == 409


async def test_docx_table_cells_keep_native_positions_and_heading():
    document = Document()
    document.add_heading("Synthetic retest", level=1)
    document.add_paragraph("Keep the same calibration revision.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Point ID"
    table.cell(0, 1).text = "Direct force (kN)"
    table.cell(1, 0).text = "Synthetic T01"
    table.cell(1, 1).text = "123"
    buffer = BytesIO()
    document.save(buffer)
    parsed = await parse_knowledge_content(buffer.getvalue(), DOCX_TYPE)
    cell = next(piece for piece in parsed.passages if piece.text == "123")
    assert cell.native_locator["kind"] == "docx_table_cell"
    assert cell.native_locator["row_index"] == 1 and cell.native_locator["column_index"] == 1
    assert cell.section == "Synthetic retest" and cell.page_number is None
    assert "python-docx=" in cell.parser_version
    assert parsed.body[cell.char_start : cell.char_end] == "123"


async def test_blank_scanned_pdf_is_explicitly_unusable_without_ocr():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="OCR"):
        await parse_knowledge_content(buffer.getvalue(), "application/pdf")


async def test_encrypted_and_invalid_documents_do_not_produce_fake_passages():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.encrypt("synthetic-test-password")
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="encrypted"):
        await parse_knowledge_content(buffer.getvalue(), "application/pdf")
    with pytest.raises(ValueError):
        await parse_knowledge_content(b"not a PDF", "application/pdf")
    with pytest.raises(ValueError):
        await parse_knowledge_content(b"\xff\xfe", "text/plain")


def test_legacy_text_is_preserved_with_no_verified_file_position():
    body = "  historical text  \n" * 500
    parsed = parse_legacy_body(body)
    assert parsed.body == body
    assert len(parsed.passages) > 1
    assert all(piece.page_number is None for piece in parsed.passages)
    assert all(piece.native_locator["source_verified"] is False for piece in parsed.passages)
    assert "".join(piece.text for piece in parsed.passages) == body
    with pytest.raises(ValueError):
        parse_legacy_body("x" * (MAX_PASSAGES * PASSAGE_CHARACTERS + 1))


async def test_passage_retrieval_ranks_matching_section_and_projects_canonical_citation(
    app, client
):
    payload = await ingest(
        app,
        client,
        (
            b"# Lubrication\n\nOil gearbox bearings.\n\n# Tendon retest\n\n"
            b"Tendon retest direct force calibration temperature.\n"
        ),
    )
    async with app.state.session_factory() as session:
        adapter = SQLToolAdapter(session)
        matches = await adapter.query_similar_failures(
            "Tendon retest direct force calibration temperature", vectorize_missing=False, limit=1
        )
        assert matches[0]["document_id"] == payload["document_id"]
        assert "direct force" in matches[0]["excerpt"]
        assert matches[0]["passage_id"]
        assert matches[0]["section"] == "Tendon retest"
        snapshot = await build_knowledge_graph_snapshot(session)
        node = next(node for node in snapshot.nodes if node.entity_id == matches[0]["passage_id"])
        assert node.node_type is NodeType.KNOWLEDGE_PASSAGE
        assert node.properties["citationUri"] == matches[0]["citation_uri"]
        assert node.properties["sourceSha256"] == payload["artifact_sha256"]
        assert node.properties["pageNumber"] is None


async def test_passage_permission_filter_precedes_ranking_and_source_download(app, client):
    async with app.state.session_factory() as session, session.begin():
        session.add(Tenant(id="passage-other-tenant", name="Synthetic other tenant"))
    payload = await ingest(
        app, client, b"Hidden tendon retest calibration details.", tenant_id="passage-other-tenant"
    )
    hidden = (
        await client.get(f"/api/v1/knowledge/documents/{payload['document_id']}/passages")
    ).json()["passages"][0]
    headers = {
        "X-WindOps-Test-Principal": "east-passage-reader",
        "X-WindOps-Test-Role": "operations_manager",
        "X-WindOps-Test-Tenant-Ids": "tenant-east-china",
        "X-WindOps-Test-Data-Scopes": "knowledge",
    }
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test", headers=headers
    ) as scoped:
        for path in (
            hidden["citation_href"],
            hidden["source_href"],
            f"/api/v1/knowledge/documents/{payload['document_id']}/passages",
        ):
            assert (await scoped.get(path)).status_code == 404
    async with app.state.session_factory() as session:
        adapter = SQLToolAdapter(
            session,
            knowledge_policy=GraphAccessPolicy.from_values(
                tenant_ids=["tenant-east-china"], data_scopes=["knowledge"]
            ),
        )
        result = await adapter.query_similar_failures(
            "Hidden tendon retest calibration details", limit=100, vectorize_missing=False
        )
        assert hidden["passage_id"] not in {row.get("passage_id") for row in result}


async def test_unrelated_passage_cannot_be_used_as_source_locator(app, client):
    first = await ingest(app, client, b"first synthetic document")
    second = await ingest(app, client, b"second synthetic document")
    other = (
        await client.get(f"/api/v1/knowledge/documents/{second['document_id']}/passages")
    ).json()["passages"][0]
    mismatch = await client.get(
        f"/api/v1/knowledge/documents/{first['document_id']}/source",
        params={"passage_id": other["passage_id"]},
    )
    assert mismatch.status_code == 404


@pytest.mark.parametrize("global_document", [False, True])
async def test_document_entity_grant_is_case_insensitive_and_does_not_expand_global_reads(
    app, client, global_document
):
    scope = {} if global_document else {"tenant_id": "tenant-east-china"}
    first = await ingest(app, client, b"Allowed synthetic source", **scope)
    sibling = await ingest(app, client, b"Sibling must stay hidden", **scope)
    headers = {
        "X-WindOps-Test-Principal": "document-entity-reader",
        "X-WindOps-Test-Role": "operations_manager",
        "X-WindOps-Test-Entity-Ids": first["document_id"].lower(),
        "X-WindOps-Test-Allow-Global": "true" if global_document else "false",
        "X-WindOps-Test-Data-Scopes": "knowledge",
    }
    if not global_document:
        headers["X-WindOps-Test-Tenant-Ids"] = "tenant-east-china"
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test", headers=headers
    ) as scoped:
        allowed = await scoped.get(f"/api/v1/knowledge/documents/{first['document_id']}/passages")
        assert allowed.status_code == 200, allowed.text
        passage = allowed.json()["passages"][0]
        assert (await scoped.get(passage["citation_href"])).status_code == 200
        assert (await scoped.get(passage["source_href"])).status_code == 200
        assert (
            await scoped.get(f"/api/v1/knowledge/documents/{sibling['document_id']}/source")
        ).status_code == 404
    async with app.state.session_factory() as session:
        policy = GraphAccessPolicy.from_values(
            tenant_ids=[] if global_document else ["tenant-east-china"],
            entity_ids=[first["document_id"].lower()],
            data_scopes=["knowledge"],
            allow_global=global_document,
        )
        adapter = SQLToolAdapter(session, knowledge_policy=policy)
        matches = await adapter.query_similar_failures(
            "synthetic source", limit=100, vectorize_missing=False
        )
        assert {row["document_id"] for row in matches} == {first["document_id"]}


async def test_passage_download_is_blocked_when_read_audit_admission_fails(app, client):
    payload = await ingest(app, client, b"synthetic audited source")

    class FailingSink:
        async def record(self, _event):
            raise ReadAuditEnqueueError("READ_AUDIT_BACKPRESSURE", "synthetic backpressure")

    app.state.read_audit_sink = FailingSink()
    response = await client.get(f"/api/v1/knowledge/documents/{payload['document_id']}/source")
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "READ_AUDIT_BACKPRESSURE"


@pytest.mark.parametrize("provider_error", [False, True])
async def test_stale_knowledge_worker_does_not_publish_passage_vectors(app, provider_error):
    parsed = await parse_knowledge_content(
        b"Synthetic immutable structural procedure.", "text/plain"
    )
    request = KnowledgeDocumentCreateRequest(
        document_id="KB-FENCED-PASSAGE",
        title="Synthetic source",
        document_type="maintenance-procedure",
        artifact_uri="minio://windops-knowledge-documents/documents/KB-FENCED-PASSAGE/source",
        artifact_sha256="a" * 64,
        content_type="text/plain",
    )
    async with app.state.session_factory() as session, session.begin():
        _, _, event_id = await create_knowledge_document(
            session,
            request,
            extracted_text=parsed.body,
            content_size_bytes=50,
            subject="synthetic-test",
            parsed=parsed,
        )

    class LeaseReplacingProvider(DeterministicTestEmbeddingProvider):
        async def embed(self, texts):
            values = await super().embed(texts)
            async with app.state.session_factory() as session, session.begin():
                event = await session.get(OutboxEvent, event_id)
                event.claim_token = str(uuid4())
            if provider_error:
                raise TimeoutError("synthetic provider failed after claim replacement")
            return values

    assert (
        await process_knowledge_document_index_event(
            app.state.session_factory,
            app.state.settings.model_copy(update={"embedding_max_retries": 0}),
            event_id,
            LeaseReplacingProvider(),
        )
        is False
    )
    async with app.state.session_factory() as session:
        rows = (
            await session.scalars(
                select(KnowledgePassage).where(KnowledgePassage.document_id == request.document_id)
            )
        ).all()
        assert len(rows) == 1 and rows[0].embedding is None
        document = await session.get(KnowledgeDocument, request.document_id)
        assert document.vectorized is False
        assert document.ingestion_status == "pending"
