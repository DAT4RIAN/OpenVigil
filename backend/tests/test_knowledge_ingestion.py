import asyncio
import hashlib
from datetime import UTC, datetime
from io import BytesIO

import httpx
import pytest
from docx import Document
from fastapi import FastAPI
from httpx import AsyncClient
from sqlalchemy import func, select

from windops_backend.agents.tools import (
    EMBEDDING_MAX_INPUT_CHARACTERS,
    DeterministicTestEmbeddingProvider,
    SQLToolAdapter,
    chunk_embedding_text,
    embed_texts_with_controls,
)
from windops_backend.knowledge_graph.domain import GraphAccessPolicy
from windops_backend.models import DomainEvent, KnowledgeDocument, Tenant
from windops_backend.schemas import AgentToolExecuteRequest, KnowledgeDocumentCreateRequest
from windops_backend.services.agent_governance import execute_agent_tool
from windops_backend.services.knowledge import (
    KNOWLEDGE_DOCUMENT_INDEX_REQUESTED,
    create_knowledge_document,
    extract_knowledge_text,
    process_knowledge_document_index_event,
)
from windops_backend.storage import InMemoryArtifactVerifier, OutboxEvent


class InvalidEmbeddingProvider:
    provider_name = "invalid-test-provider"
    model_name = "wrong-dimensions"
    production_ready = False

    async def embed(self, texts: list[str]) -> list[list[float]]:
        return [[1.0] for _text in texts]


class RecordingEmbeddingProvider(DeterministicTestEmbeddingProvider):
    def __init__(self) -> None:
        self.calls: list[list[str]] = []

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls.append(list(texts))
        return await super().embed(texts)


class RetryOnceEmbeddingProvider(RecordingEmbeddingProvider):
    def __init__(self) -> None:
        super().__init__()
        self.failed = False

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls.append(list(texts))
        if not self.failed:
            self.failed = True
            raise TimeoutError("controlled embedding timeout")
        return await DeterministicTestEmbeddingProvider.embed(self, texts)


class CancelledEmbeddingProvider:
    provider_name = "cancelled-test-provider"
    model_name = "cancelled-test-model"
    production_ready = False

    async def embed(self, texts: list[str]) -> list[list[float]]:
        del texts
        raise asyncio.CancelledError


def test_knowledge_embedding_chunks_are_bounded_and_overlap() -> None:
    chunks = chunk_embedding_text("Main bearing", "x" * 30_000)
    assert len(chunks) > 1
    assert all(len(chunk) <= EMBEDDING_MAX_INPUT_CHARACTERS for chunk in chunks)
    prefix_length = len("Main bearing\n")
    assert chunks[0][-1_000:] == chunks[1][prefix_length : prefix_length + 1_000]


async def test_embedding_batches_have_timeout_retry_and_batch_limits(app: FastAPI) -> None:
    provider = RetryOnceEmbeddingProvider()
    vectors = await embed_texts_with_controls(
        provider,
        [f"controlled text {index}" for index in range(17)],
        settings=app.state.settings,
    )
    assert len(vectors) == 17
    assert [len(call) for call in provider.calls] == [16, 16, 1]


async def test_large_knowledge_document_is_indexed_as_bounded_batches(app: FastAPI) -> None:
    request = KnowledgeDocumentCreateRequest(
        document_id="KB-LONG-001",
        title="Long controlled manual",
        document_type="maintenance-procedure",
        artifact_uri="minio://windops-knowledge-documents/documents/KB-LONG-001/manual.txt",
        artifact_sha256="a" * 64,
        content_type="text/plain",
    )
    async with app.state.session_factory() as session, session.begin():
        _document, _replayed, event_id = await create_knowledge_document(
            session,
            request,
            extracted_text="bearing lubrication procedure " * 6_000,
            content_size_bytes=180_000,
            subject="knowledge-manager",
        )
    provider = RecordingEmbeddingProvider()
    assert event_id is not None
    await process_knowledge_document_index_event(
        app.state.session_factory,
        app.state.settings,
        event_id,
        provider,
    )
    assert len(provider.calls) > 1
    assert max(len(call) for call in provider.calls) <= app.state.settings.embedding_batch_size
    assert all(
        len(text) <= EMBEDDING_MAX_INPUT_CHARACTERS for call in provider.calls for text in call
    )
    async with app.state.session_factory() as session:
        indexed = await session.get(KnowledgeDocument, request.document_id)
        assert indexed is not None and indexed.vectorized is True


async def test_similarity_query_does_not_auto_index_the_visible_corpus(app: FastAPI) -> None:
    async with app.state.session_factory() as session, session.begin():
        for index in range(12):
            session.add(
                KnowledgeDocument(
                    id=f"KB-PENDING-{index:02d}",
                    title=f"Pending manual {index}",
                    document_type="maintenance-procedure",
                    body="pending body",
                    citation_uri=f"windops://knowledge/documents/KB-PENDING-{index:02d}",
                    vectorized=False,
                    created_by="knowledge-manager",
                )
            )
    provider = RecordingEmbeddingProvider()
    async with app.state.session_factory() as session:
        adapter = SQLToolAdapter(session, embedding_provider=provider)
        await adapter.query_similar_failures("unindexed query", limit=5, vectorize_missing=False)
    assert provider.calls == [["unindexed query"]]


async def test_cancelled_knowledge_indexing_stays_pending_for_lease_recovery(
    app: FastAPI,
) -> None:
    request = KnowledgeDocumentCreateRequest(
        document_id="KB-CANCELLED-001",
        title="Cancelled indexing test",
        document_type="maintenance-procedure",
        artifact_uri="minio://windops-knowledge-documents/documents/KB-CANCELLED-001/test.txt",
        artifact_sha256="c" * 64,
        content_type="text/plain",
    )
    async with app.state.session_factory() as session, session.begin():
        _document, _replayed, event_id = await create_knowledge_document(
            session,
            request,
            extracted_text="Controlled body for cancellation recovery.",
            content_size_bytes=42,
            subject="knowledge-manager",
        )
    assert event_id is not None
    with pytest.raises(asyncio.CancelledError):
        await process_knowledge_document_index_event(
            app.state.session_factory,
            app.state.settings,
            event_id,
            CancelledEmbeddingProvider(),
        )
    async with app.state.session_factory() as session:
        document = await session.get(KnowledgeDocument, request.document_id)
        event = await session.get(OutboxEvent, event_id)
        assert document is not None and document.ingestion_status == "pending"
        assert document.vectorized is False
        assert event is not None and event.status == "processing"


async def test_text_knowledge_artifact_is_verified_indexed_and_replay_safe(
    app: FastAPI, client: AsyncClient
) -> None:
    content = b"# Main bearing inspection\n\nVerify lubrication, vibration and temperature.\n"
    digest = hashlib.sha256(content).hexdigest()
    artifact_uri = "minio://windops-knowledge-documents/documents/KB-PROD-001/manual.md"
    verifier = app.state.artifact_verifier
    assert isinstance(verifier, InMemoryArtifactVerifier)
    assert verifier.register_object(artifact_uri, content, "text/markdown") == digest
    payload = {
        "document_id": "KB-PROD-001",
        "title": "Main bearing controlled inspection",
        "document_type": "maintenance-procedure",
        "document_version": "2026.08",
        "artifact_uri": artifact_uri,
        "artifact_sha256": digest,
        "content_type": "text/markdown",
        "metadata": {
            "manufacturer": "Goldwind",
            "equipment": "main-bearing",
            "tags": ["inspection", "bearing"],
        },
    }
    command_headers = {"Idempotency-Key": "knowledge-document-prod-001"}
    created = await client.post(
        "/api/v1/knowledge/documents", headers=command_headers, json=payload
    )
    assert created.status_code == 202
    assert created.headers["Idempotency-Replayed"] == "false"
    body = created.json()
    assert body["ingestion_status"] == "indexed"
    assert body["vectorized"] is True
    assert body["embedding_provider"] == "deterministic_test"
    assert body["artifact_sha256"] == digest

    collection = await client.get(
        "/api/v1/knowledge/documents", params={"q": "bearing", "ingestion_status": "indexed"}
    )
    assert collection.status_code == 200
    assert collection.json()["documents"][0]["document_id"] == "KB-PROD-001"
    detail = await client.get("/api/v1/knowledge/documents/KB-PROD-001")
    assert detail.status_code == 200
    assert "Verify lubrication" in detail.json()["body"]

    replay = await client.post("/api/v1/knowledge/documents", headers=command_headers, json=payload)
    assert replay.status_code == 202
    assert replay.headers["Idempotency-Replayed"] == "true"
    assert replay.json() == created.json()

    async with app.state.session_factory() as session:
        outbox = await session.scalar(
            select(OutboxEvent).where(
                OutboxEvent.event_type == KNOWLEDGE_DOCUMENT_INDEX_REQUESTED,
                OutboxEvent.aggregate_id == "KB-PROD-001",
            )
        )
        assert outbox is not None and outbox.status == "succeeded"
        assert (
            await session.scalar(
                select(func.count())
                .select_from(DomainEvent)
                .where(DomainEvent.event_type == "knowledge.document.indexed")
            )
            == 1
        )

    changed = b"different controlled content"
    changed_uri = "minio://windops-knowledge-documents/documents/KB-PROD-001/changed.md"
    changed_hash = verifier.register_object(changed_uri, changed, "text/markdown")
    conflict = await client.post(
        "/api/v1/knowledge/documents",
        headers=command_headers,
        json={
            **payload,
            "artifact_uri": changed_uri,
            "artifact_sha256": changed_hash,
        },
    )
    assert conflict.status_code == 409
    assert conflict.json()["error"]["code"] == "IDEMPOTENCY_KEY_REUSED"


async def test_knowledge_upload_and_ingest_require_manager_role(
    client: AsyncClient,
) -> None:
    forbidden_headers = {
        "X-WindOps-Test-Principal": "field-tech",
        "X-WindOps-Test-Role": "field_technician",
        "Idempotency-Key": "knowledge-upload-forbidden-001",
    }
    upload = await client.post(
        "/api/v1/knowledge/documents/uploads/presign",
        headers=forbidden_headers,
        json={
            "document_id": "KB-FORBIDDEN-001",
            "file_name": "manual.md",
            "content_type": "text/markdown",
            "artifact_sha256": "a" * 64,
        },
    )
    assert upload.status_code == 403


async def test_knowledge_documents_and_agent_reads_are_tenant_scoped(
    app: FastAPI,
) -> None:
    async with app.state.session_factory() as session, session.begin():
        session.add(Tenant(id="tenant-west", name="West Wind Operations", status="active"))
        session.add(
            KnowledgeDocument(
                id="KB-WEST-001",
                tenant_id="tenant-west",
                data_scope="knowledge",
                title="West tenant restricted manual",
                document_type="maintenance-procedure",
                body="West tenant body must not cross the tenant boundary.",
                citation_uri="windops://knowledge/documents/KB-WEST-001",
                vectorized=False,
                created_by="west-manager",
            )
        )

    scoped_headers = {
        "X-WindOps-Test-Principal": "east-reader",
        "X-WindOps-Test-Role": "operations_manager",
        "X-WindOps-Test-Tenant-Ids": "tenant-east-china",
    }
    scoped_client = httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test", headers=scoped_headers
    )
    async with scoped_client:
        collection = await scoped_client.get("/api/v1/knowledge/documents")
        assert collection.status_code == 200
        document_ids = {row["document_id"] for row in collection.json()["documents"]}
        assert "KB-MB-GW165-001" in document_ids
        assert "KB-WEST-001" not in document_ids

        hidden = await scoped_client.get("/api/v1/knowledge/documents/KB-WEST-001")
        assert hidden.status_code == 404
        assert hidden.json()["error"]["code"] == "NOT_FOUND"

    async with app.state.session_factory() as session:
        adapter = SQLToolAdapter(
            session,
            knowledge_policy=GraphAccessPolicy.from_values(tenant_ids=["tenant-east-china"]),
        )
        visible = await adapter.visible_knowledge_documents()
        assert "KB-WEST-001" not in {document.id for document in visible}
        matches = await adapter.query_similar_failures("restricted manual", limit=20)
        assert "KB-WEST-001" not in {row["document_id"] for row in matches}

        invocation, replayed = await execute_agent_tool(
            session,
            AgentToolExecuteRequest(
                tool="query_manual",
                args={"query": "manual", "limit": 20},
                agent_id="knowledge_agent",
                idempotency_key="knowledge-scope-test-001",
                persist=False,
            ),
            subject="east-reader",
            knowledge_policy=GraphAccessPolicy.from_values(tenant_ids=["tenant-east-china"]),
        )
        assert replayed is False
        assert invocation.status == "succeeded"
        assert "KB-WEST-001" not in {row["document_id"] for row in invocation.result["rows"]}


@pytest.mark.parametrize(
    "scope_alias",
    ["knowledge", "knowledges", "knowledge_document", "knowledge_documents"],
)
async def test_knowledge_scope_aliases_cover_write_list_detail_and_vector_reads(
    app: FastAPI,
    scope_alias: str,
) -> None:
    alias_id = scope_alias.replace("_", "-").upper()
    document_id = f"KB-ALIAS-{alias_id}"
    content = f"Alias controlled {scope_alias} manual for bearing inspection.".encode()
    artifact_uri = f"minio://windops-knowledge-documents/documents/{document_id}/manual.txt"
    verifier = app.state.artifact_verifier
    assert isinstance(verifier, InMemoryArtifactVerifier)
    digest = verifier.register_object(artifact_uri, content, "text/plain")
    headers = {
        "X-WindOps-Test-Principal": f"{scope_alias}-reader",
        "X-WindOps-Test-Role": "operations_manager",
        "X-WindOps-Test-Tenant-Ids": "tenant-east-china",
        "X-WindOps-Test-Data-Scopes": scope_alias,
        "Idempotency-Key": f"knowledge-scope-{scope_alias}-001",
    }
    payload = {
        "document_id": document_id,
        "title": f"Alias controlled {scope_alias} manual",
        "document_type": "maintenance-procedure",
        "document_version": "1",
        "artifact_uri": artifact_uri,
        "artifact_sha256": digest,
        "content_type": "text/plain",
        "tenant_id": "tenant-east-china",
        "data_scope": "knowledge",
    }

    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://test",
        headers=headers,
    ) as scoped_client:
        created = await scoped_client.post("/api/v1/knowledge/documents", json=payload)
        assert created.status_code == 202, created.text

        collection = await scoped_client.get("/api/v1/knowledge/documents")
        assert collection.status_code == 200, collection.text
        assert document_id in {row["document_id"] for row in collection.json()["documents"]}

        detail = await scoped_client.get(f"/api/v1/knowledge/documents/{document_id}")
        assert detail.status_code == 200, detail.text
        assert scope_alias in detail.json()["body"]

    async with app.state.session_factory() as session:
        adapter = SQLToolAdapter(
            session,
            knowledge_policy=GraphAccessPolicy.from_values(
                tenant_ids=["tenant-east-china"], data_scopes=[scope_alias]
            ),
        )
        matches = await adapter.query_similar_failures(
            f"Alias controlled {scope_alias} manual", limit=100
        )
        assert document_id in {row["document_id"] for row in matches}


async def test_docx_text_extraction_and_failed_embedding_are_explicit(
    app: FastAPI,
) -> None:
    docx = Document()
    docx.add_heading("Gearbox cooling", level=1)
    docx.add_paragraph("Inspect the cooling circuit before return to service.")
    buffer = BytesIO()
    docx.save(buffer)
    text = await extract_knowledge_text(
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "Inspect the cooling circuit" in text

    request = KnowledgeDocumentCreateRequest(
        document_id="KB-FAILED-001",
        title="Embedding failure test",
        document_type="maintenance-procedure",
        artifact_uri="minio://windops-knowledge-documents/documents/KB-FAILED-001/test.txt",
        artifact_sha256="f" * 64,
        content_type="text/plain",
    )
    async with app.state.session_factory() as session, session.begin():
        document, replayed, event_id = await create_knowledge_document(
            session,
            request,
            extracted_text="Controlled body for indexing.",
            content_size_bytes=29,
            subject="knowledge-manager",
        )
        assert replayed is False and event_id is not None
        assert document.created_at <= datetime.now(UTC)
    with pytest.raises(RuntimeError, match="expected 1536"):
        await process_knowledge_document_index_event(
            app.state.session_factory,
            app.state.settings,
            event_id,
            InvalidEmbeddingProvider(),
        )
    async with app.state.session_factory() as session:
        failed = await session.get(KnowledgeDocument, "KB-FAILED-001")
        event = await session.get(OutboxEvent, event_id)
        assert failed is not None and failed.ingestion_status == "failed"
        assert failed.vectorized is False
        assert event is not None and event.status == "failed"


async def test_malformed_document_parser_errors_are_normalized(app: FastAPI) -> None:
    with pytest.raises(ValueError, match="could not be parsed"):
        await extract_knowledge_text(b"not a PDF", "application/pdf")
